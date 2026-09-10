import asyncio
import io
import os

from docx import Document


def make_test_docx() -> bytes:
    """在内存里生成一份测试简历 DOCX"""
    doc = Document()
    doc.add_heading("张三 - Python 后端开发工程师", level=1)
    doc.add_paragraph("学历：本科    工作年限：3年")
    doc.add_paragraph("联系方式：zhangsan@example.com | 138-xxxx-xxxx")

    doc.add_heading("技能", level=2)
    doc.add_paragraph("Python, FastAPI, Django, MySQL, Redis, Docker, Linux, Git")

    doc.add_heading("工作经历", level=2)
    doc.add_paragraph("2021.07 - 2024.07  某互联网公司  后端开发工程师")
    doc.add_paragraph("- 负责核心业务系统后端开发，使用 Python + FastAPI 搭建微服务")
    doc.add_paragraph("- 优化数据库查询，引入 Redis 缓存，接口响应时间降低 60%")
    doc.add_paragraph("- 参与爬虫系统开发，抓取百万级数据并清洗入库")

    doc.add_heading("项目经验", level=2)
    doc.add_paragraph("- 智能招聘系统：基于 LangChain 实现简历解析和岗位匹配")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def main():
    # 准备测试数据
    docx_bytes = make_test_docx()
    print(f"已生成测试 DOCX，{len(docx_bytes)} 字节")

    # 保存一份到 data/ 目录供后续测试用
    os.makedirs("data", exist_ok=True)
    test_path = "data/test_resume.docx"
    with open(test_path, "wb") as f:
        f.write(docx_bytes)
    print(f"已保存到 {test_path}")

    # 直接用 httpx 调接口（不启服务，用 ASGI 直连）
    from httpx import ASGITransport, AsyncClient
    from app.main import app

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 上传
        files = {"file": ("test_resume.docx", docx_bytes,
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = await client.post("/api/v1/resumes/upload", files=files)
        print(f"\n上传返回 [{r.status_code}]:")
        print(r.json())

        if r.status_code == 200:
            rid = r.json()["id"]
            # 再查一次
            r2 = await client.get(f"/api/v1/resumes/{rid}")
            print(f"\n查询返回 [{r2.status_code}]:")
            print(r2.json())


if __name__ == "__main__":
    asyncio.run(main())