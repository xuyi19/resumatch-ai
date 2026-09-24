"""M20：报告导出（A2）+ 重诊溯源 parent_task_id（A3）"""
import io

from docx import Document
from httpx import ASGITransport, AsyncClient
import pytest

from app.main import app


def _fake_result(overall=72):
    """与 _finalize 组装结构一致的诊断 result"""
    return {
        "keyword": "后端开发工程师",
        "city": "上海",
        "jd_text": "负责后端服务开发，要求熟悉 Python 与 FastAPI，" + "细节" * 20,
        "diagnosis": {
            "scores": {
                "overall": overall,
                "completeness": {"score": 80, "comment": "信息完整", "evidence_ids": []},
                "quantification": {"score": 65, "comment": "量化一般", "evidence_ids": []},
                "star_structure": {"score": 70, "comment": "STAR 尚可", "evidence_ids": []},
                "skill_match": {"score": 75, "comment": "技能匹配", "evidence_ids": []},
                "achievement": {"score": 60, "comment": "亮点不足", "evidence_ids": []},
                "readability": {"score": 85, "comment": "可读性好", "evidence_ids": []},
            },
            "gaps": [
                {"dimension": "量化", "description": "缺少量化数据",
                 "severity": "high", "is_inferred": False,
                 "evidence": [{"id": "R1", "text": "负责开发系统"}]},
                {"dimension": "技能", "description": "缺少中间件经验",
                 "severity": "medium", "is_inferred": False, "evidence": []},
            ],
            "suggestions": [
                {"target": "项目经历", "original": "负责开发系统",
                 "rewritten": "主导开发 XX 系统，QPS 提升 40%",
                 "reason": "增加量化", "evidence_ids": ["R1"]},
            ],
            "overall_advice": "整体补充量化数据",
        },
    }


async def _seed_success_record(task_id="task_report1", parent=None, overall=72):
    from sqlalchemy import delete

    from app.core.db import AsyncSessionLocal
    from app.models.entities import DiagnosisRecord

    # 幂等：test.db 跨 pytest 运行持久化，先清同 id 旧记录防唯一约束冲突
    async with AsyncSessionLocal() as session:
        await session.execute(
            delete(DiagnosisRecord).where(DiagnosisRecord.task_id == task_id)
        )
        session.add(DiagnosisRecord(
            task_id=task_id,
            resume_id=None,
            keyword="后端开发工程师",
            resume_name="测试简历",
            owner_id="local",
            status="success",
            result=_fake_result(overall),
            parent_task_id=parent,
        ))
        await session.commit()


@pytest.mark.asyncio
async def test_export_report_docx():
    """成功记录可导出报告 docx：内容含评分表格与建议段落"""
    await _seed_success_record("task_export_ok")
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.post("/api/v1/history/task_export_ok/export-report", json={})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        doc = Document(io.BytesIO(r.content))
        texts = [p.text for p in doc.paragraphs]
        assert any("简历诊断报告" in t for t in texts)
        assert any("综合评分" in t for t in texts)
        # 表格存在且包含六维行
        assert doc.tables and len(doc.tables[0].rows) == 7
        assert any("改写建议" in t for t in texts)
        # 导出历史落库（template=report）
        r2 = await client.get("/api/v1/resumes/export-history")
        templates = [i["template"] for i in r2.json()["items"]]
        assert "report" in templates


@pytest.mark.asyncio
async def test_export_report_rejects_incomplete():
    """无完整结果的记录拒绝导出"""
    from sqlalchemy import delete

    from app.core.db import AsyncSessionLocal
    from app.models.entities import DiagnosisRecord

    async with AsyncSessionLocal() as session:
        await session.execute(
            delete(DiagnosisRecord).where(DiagnosisRecord.task_id == "task_failed1")
        )
        session.add(DiagnosisRecord(
            task_id="task_failed1", keyword="岗位", owner_id="local",
            status="failed", error="boom",
        ))
        await session.commit()
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.post("/api/v1/history/task_failed1/export-report", json={})
        assert r.status_code == 409


@pytest.mark.asyncio
async def test_history_list_contains_score_and_parent():
    """历史列表带 score 与 parent_task_id（D4/A3 前端依赖）"""
    await _seed_success_record("task_list_score", parent="task_parent_x", overall=88)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.get("/api/v1/history", params={"limit": 50})
        assert r.status_code == 200
        items = {i["task_id"]: i for i in r.json()["items"]}
        assert items["task_list_score"]["score"] == 88
        assert items["task_list_score"]["parent_task_id"] == "task_parent_x"


@pytest.mark.asyncio
async def test_history_detail_contains_resume_id():
    """历史详情带 resume_id（前端「用当前简历重新诊断」依赖）"""
    from sqlalchemy import delete

    from app.core.db import AsyncSessionLocal
    from app.models.entities import DiagnosisRecord

    async with AsyncSessionLocal() as session:
        await session.execute(
            delete(DiagnosisRecord).where(DiagnosisRecord.task_id == "task_detail1")
        )
        session.add(DiagnosisRecord(
            task_id="task_detail1", resume_id=42, keyword="岗位",
            owner_id="local", status="success", result=_fake_result(),
        ))
        await session.commit()
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.get("/api/v1/history/task_detail1")
        assert r.status_code == 200
        assert r.json()["resume_id"] == 42


@pytest.mark.asyncio
async def test_report_docx_generates():
    """generate_report_docx 纯函数：最小 result 不抛异常"""
    from app.utils.report_docx import generate_report_docx

    data = generate_report_docx({"keyword": "岗位", "diagnosis": {"scores": {}}})
    assert data[:2] == b"PK"  # docx 是 zip 包
