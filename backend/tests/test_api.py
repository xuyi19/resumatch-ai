"""API 接口集成测试"""
import io

from docx import Document
from httpx import ASGITransport, AsyncClient
import pytest

from app.main import app


def make_test_docx() -> bytes:
    doc = Document()
    doc.add_heading("测试简历", level=1)
    doc.add_paragraph("学历：本科    工作年限：3年")
    doc.add_paragraph("技能：Python, FastAPI, MySQL")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_upload_resume():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        docx_bytes = make_test_docx()
        files = {
            "file": ("test.docx", docx_bytes,
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        r = await client.post("/api/v1/resumes/upload", files=files)
        assert r.status_code == 200
        data = r.json()
        assert "id" in data
        assert data["text_length"] > 0
        print(f"\n✓ 简历上传成功 id={data['id']} 长度={data['text_length']}")


@pytest.mark.asyncio
async def test_upload_invalid_file():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        files = {"file": ("test.xyz", b"invalid content", "text/plain")}
        r = await client.post("/api/v1/resumes/upload", files=files)
        assert r.status_code == 400


@pytest.mark.asyncio
async def test_live_analyze_validation():
    """JD 文本过短应返回 400"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.post("/api/v1/live/analyze", json={
            "resume_id": 999999,
            "jd_text": "太短",
        })
        assert r.status_code == 400


@pytest.mark.asyncio
async def test_templates_endpoint():
    """模板目录接口应返回全部模板"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.get("/api/v1/resumes/templates")
        assert r.status_code == 200
        items = r.json()["items"]
        assert len(items) >= 10
        assert all("id" in t and "name" in t for t in items)


def _make_png() -> bytes:
    """最小合法 PNG（1x1 白色像素）"""
    import struct
    import zlib

    def chunk(tag, data):
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\xff\xff")
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", idat)
            + chunk(b"IEND", b""))


@pytest.mark.asyncio
async def test_photo_upload_and_export():
    """上传证件照 → 带照片导出 Word"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.post(
            "/api/v1/resumes/photo",
            files={"file": ("photo.png", _make_png(), "image/png")},
        )
        assert r.status_code == 200
        photo_id = r.json()["photo_id"]
        assert len(photo_id) == 32

        # 照片可回读
        r2 = await client.get(f"/api/v1/resumes/photo/{photo_id}")
        assert r2.status_code == 200
        assert r2.headers["content-type"] == "image/png"

        # 拒绝非图片
        r3 = await client.post(
            "/api/v1/resumes/photo",
            files={"file": ("a.txt", b"not an image", "text/plain")},
        )
        assert r3.status_code == 400

        # 带照片导出
        resume_data = {
            "name": "张三",
            "job_intention": "后端开发",
            "contact": "电话：13800000000",
            "summary": "三年后端经验",
            "education": ["2017-2021 某大学 计算机 本科"],
            "experience": ["2021-至今 某公司 后端开发"],
            "projects": [],
            "skills": ["Python"],
            "certificates": [],
        }
        r4 = await client.post("/api/v1/resumes/export-docx", json={
            "optimized_resume": resume_data,
            "template": "sidebar",
            "photo_id": photo_id,
        })
        assert r4.status_code == 200
        assert r4.content[:2] == b"PK"  # docx 是 zip 包

        # 清理
        r5 = await client.delete(f"/api/v1/resumes/photo/{photo_id}")
        assert r5.status_code == 200


@pytest.mark.asyncio
async def test_owner_isolation(monkeypatch):
    """web 模式下不同会话的数据互相隔离"""
    from app.core.config import settings

    monkeypatch.setattr(settings, "APP_MODE", "web")
    transport = ASGITransport(app=app)

    async with AsyncClient(transport=transport, base_url="http://test") as client_a:
        files = {
            "file": ("a.docx", make_test_docx(),
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        r = await client_a.post("/api/v1/resumes/upload", files=files)
        assert r.status_code == 200
        resume_id = r.json()["id"]

        # 会话 A 自己可见
        r2 = await client_a.get(f"/api/v1/resumes/{resume_id}")
        assert r2.status_code == 200

        # 会话 B 看不到 A 的简历，也看不到 A 的历史记录
        async with AsyncClient(transport=transport, base_url="http://test") as client_b:
            r3 = await client_b.get(f"/api/v1/resumes/{resume_id}")
            assert r3.status_code == 404
            r4 = await client_b.get("/api/v1/history")
            task_ids = {item["task_id"] for item in r4.json()["items"]}
            r5 = await client_a.get("/api/v1/history")
            a_task_ids = {item["task_id"] for item in r5.json()["items"]}
            assert not (a_task_ids & task_ids)


@pytest.mark.asyncio
async def test_upload_resume_text():
    """A1 粘贴文本简历：直存入库，跳过文件解析"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        text = "张三\n本科，3 年 Python 后端经验，熟悉 FastAPI 与 MySQL，做过电商订单系统。" * 3
        r = await client.post("/api/v1/resumes/upload-text", json={
            "filename": "粘贴的简历.txt",
            "text": text,
        })
        assert r.status_code == 200
        data = r.json()
        assert data["filename"] == "粘贴的简历.txt"
        assert data["text_length"] == len(text.strip())

        # 文本过短应拒绝
        r2 = await client.post("/api/v1/resumes/upload-text", json={
            "filename": "短.txt", "text": "太短了",
        })
        assert r2.status_code == 400


@pytest.mark.asyncio
async def test_delete_resumes_batch():
    """F1 批量删除：多份简历一次删除（owner 隔离）；回归：Resume 未 import 曾致 500"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        ids = []
        for name, body in (
            ("批量删除-1.txt", "李四\n后端工程师，熟悉 Redis 与 Kafka，做过日志平台。" * 3),
            ("批量删除-2.txt", "王五\n前端工程师，熟悉 Vue3 与 TypeScript，做过组件库。" * 3),
        ):
            r = await client.post("/api/v1/resumes/upload-text", json={
                "filename": name, "text": body,
            })
            assert r.status_code == 200
            ids.append(r.json()["id"])

        r = await client.post("/api/v1/resumes/delete-batch", json={"ids": ids})
        assert r.status_code == 200, r.text
        assert r.json()["deleted"] == 2

        # 列表中确认消失
        r2 = await client.get("/api/v1/resumes/list")
        remain = {x["id"] for x in r2.json()["items"]}
        assert not (set(ids) & remain)

        # 不存在的 id：幂等返回 0；空列表直接 0
        r3 = await client.post("/api/v1/resumes/delete-batch", json={"ids": ids})
        assert r3.status_code == 200 and r3.json()["deleted"] == 0
        r4 = await client.post("/api/v1/resumes/delete-batch", json={"ids": []})
        assert r4.status_code == 200 and r4.json()["deleted"] == 0
