"""M63 求职周报测试：六源聚合 / 环比 delta / 空周兜底 / offset 边界 / Word 导出。

uuid 会话（X-Session-Id 头）做 owner 隔离（conftest 的 test.db 跨用例共享，
uuid owner 同时隔离跨 pytest 运行残留）；周窗口内的时间 API 造不出来，
直接塞库并显式指定 created_at / apply_at（naive datetime，同库内口径）。
"""

import io
from datetime import date, datetime, timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient

from app.core.config import settings
from app.core.db import AsyncSessionLocal
from app.main import app
from app.models.entities import Conversation, DiagnosisRecord, Offer, SavedJob

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
async def client(monkeypatch):
    monkeypatch.setattr(settings, "APP_MODE", "web")
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


def _sid() -> dict:
    return {"X-Session-Id": f"wr-{uuid4().hex[:12]}"}


def _monday(offset_weeks: int = 0) -> date:
    """本地时区本周 ISO 周一（可再往前推 N 周）——与端点聚合口径一致。"""
    today = date.today()
    return today - timedelta(days=today.weekday()) - timedelta(weeks=offset_weeks)


def _dt(d: date, hour: int = 10, minute: int = 0) -> datetime:
    return datetime(d.year, d.month, d.day, hour, minute)


async def _seed(owner_id: str, rows: list) -> None:
    async with AsyncSessionLocal() as session:
        for row in rows:
            session.add(row)
        await session.commit()


def _job(owner, title, **kw) -> SavedJob:
    return SavedJob(owner_id=owner, title=title, jd_text="周报测试岗位描述。", **kw)


def _conv(owner, type_, created_at) -> Conversation:
    return Conversation(owner_id=owner, task_id=f"wr-{type_}-{uuid4().hex[:8]}",
                        type=type_, created_at=created_at)


def _diag(owner, created_at) -> DiagnosisRecord:
    return DiagnosisRecord(owner_id=owner, task_id=f"wr-diag-{uuid4().hex[:8]}",
                           keyword="Java 后端", status="success", created_at=created_at)


# ---------------------------------------------------------------- 聚合与环比


@pytest.mark.asyncio
async def test_weekly_aggregate_counts(client):
    """本周六源计数 + 活跃天数（日期并集去重）；上周同口径进 prev。"""
    hd = _sid()
    owner = hd["X-Session-Id"]
    mon, last = _monday(0), _monday(1)
    await _seed(owner, [
        _job(owner, "本周投递岗", apply_at=_dt(mon + timedelta(days=1)),
             created_at=_dt(mon)),  # 本周：投递 + 入库（不同天）
        _conv(owner, "interview", _dt(mon + timedelta(days=2))),
        _conv(owner, "quiz", _dt(mon + timedelta(days=2))),  # 与面试同天
        _diag(owner, _dt(mon + timedelta(days=3))),
        Offer(owner_id=owner, company="A 科技", title="后端工程师",
              monthly_salary=20, created_at=_dt(mon + timedelta(days=4))),
        _job(owner, "上周投递岗", apply_at=_dt(last + timedelta(days=1)),
             created_at=_dt(last)),  # 上周：投递 + 入库
    ])

    r = await client.get("/api/v1/history/weekly-report", headers=hd)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["week_start"] == mon.strftime("%m-%d")
    assert data["week_end"] == (mon + timedelta(days=6)).strftime("%m-%d")

    m = data["metrics"]
    assert m["applications"] == {"value": 1, "prev": 1, "delta": 0}
    assert m["jobs_added"] == {"value": 1, "prev": 1, "delta": 0}
    assert m["interviews"] == {"value": 1, "prev": 0, "delta": 1}
    assert m["quizzes"] == {"value": 1, "prev": 0, "delta": 1}
    assert m["diagnoses"] == {"value": 1, "prev": 0, "delta": 1}
    assert m["offers"] == {"value": 1, "prev": 0, "delta": 1}
    # 活跃天 = 本周 5 个不同日期（周一~周五，面试+刷题同天去重）+ 上周 2 天
    assert data["active_days"] == {"value": 5, "prev": 2, "delta": 3}

    # 模板亮点按优先级取前 4 条；投递环比持平（delta=0 且 prev>0）
    assert len(data["highlights"]) == 4
    assert data["highlights"][0] == "📮 本周投递 1 次（与上周持平）"
    assert data["highlights"][1] == "🎤 完成 1 场模拟面试"
    assert data["highlights"][2] == "🧠 刷题 1 场"
    assert data["highlights"][3] == "📄 诊断 1 次"


@pytest.mark.asyncio
async def test_weekly_deltas(client):
    """上周 1 条投递、本周 2 条 → applications delta=+1，亮点带「环比 +1」。"""
    hd = _sid()
    owner = hd["X-Session-Id"]
    mon, last = _monday(0), _monday(1)
    await _seed(owner, [
        _job(owner, "本周投递一", apply_at=_dt(mon), created_at=_dt(mon)),
        _job(owner, "本周投递二", apply_at=_dt(mon + timedelta(days=2)),
             created_at=_dt(mon + timedelta(days=1))),
        _job(owner, "上周投递", apply_at=_dt(last + timedelta(days=3)), created_at=_dt(last)),
    ])

    data = (await client.get("/api/v1/history/weekly-report", headers=hd)).json()
    assert data["metrics"]["applications"] == {"value": 2, "prev": 1, "delta": 1}
    assert data["highlights"][0] == "📮 本周投递 2 次（环比 +1）"


@pytest.mark.asyncio
async def test_weekly_empty_fallback(client):
    """无数据 owner：全 0 + 只有一条兜底文案（接口成功即使全零也可见）。"""
    hd = _sid()
    r = await client.get("/api/v1/history/weekly-report", headers=hd)
    assert r.status_code == 200
    data = r.json()
    assert all(v == {"value": 0, "prev": 0, "delta": 0} for v in data["metrics"].values())
    assert data["active_days"] == {"value": 0, "prev": 0, "delta": 0}
    assert data["highlights"] == ["本周还没有记录，从投出一份简历开始吧"]


@pytest.mark.asyncio
async def test_weekly_offset_bounds(client):
    """week_offset > 8 → 422；= 8 可用。"""
    hd = _sid()
    r9 = await client.get("/api/v1/history/weekly-report",
                          params={"week_offset": 9}, headers=hd)
    assert r9.status_code == 422
    r8 = await client.get("/api/v1/history/weekly-report",
                          params={"week_offset": 8}, headers=hd)
    assert r8.status_code == 200, r8.text


@pytest.mark.asyncio
async def test_weekly_window_edges(client):
    """窗口左闭右开：上周日 23:59 与下周一 00:00 的记录均不计入本周。

    跨月 / 跨年由 timedelta 天数算术天然保证（datetime 无月概念），无需固定日期。
    """
    hd = _sid()
    owner = hd["X-Session-Id"]
    mon = _monday(0)
    await _seed(owner, [
        _job(owner, "窗口内", apply_at=_dt(mon), created_at=_dt(mon)),
        _job(owner, "上周日边界", created_at=_dt(mon - timedelta(days=1), 23, 59)),
        _conv(owner, "interview", _dt(mon + timedelta(days=7), 0, 0)),  # 下周一 00:00
    ])

    data = (await client.get("/api/v1/history/weekly-report", headers=hd)).json()
    assert data["metrics"]["jobs_added"]["value"] == 1
    assert data["metrics"]["applications"]["value"] == 1
    assert data["metrics"]["interviews"]["value"] == 0


# ---------------------------------------------------------------- Word 导出


@pytest.mark.asyncio
async def test_weekly_export_blob(client):
    """blob 模式：docx 字节流非空、正文含标题与结尾文案、表格 8 行；落 export_history。"""
    hd = _sid()
    owner = hd["X-Session-Id"]
    mon = _monday(0)
    await _seed(owner, [_job(owner, "导出测试岗", apply_at=_dt(mon), created_at=_dt(mon))])

    r = await client.post("/api/v1/history/weekly-export",
                          json={"week_offset": 0}, headers=hd)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(DOCX_MIME)
    assert len(r.content) > 0

    doc = Document(io.BytesIO(r.content))
    texts = [p.text for p in doc.paragraphs]
    assert any("求职周报" in t for t in texts)
    assert any("数据来自本机" in t for t in texts)
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 8  # 表头 + 六指标 + 活跃天数
    assert doc.tables[0].rows[1].cells[0].text == "投递次数"

    hist = (await client.get("/api/v1/resumes/export-history", headers=hd)).json()["items"]
    assert any(i["template"] == "weekly" and "求职周报" in i["filename"] for i in hist)


@pytest.mark.asyncio
async def test_weekly_export_save_path(client, monkeypatch, tmp_path):
    """桌面态 save_path 直写：落 {export_root}/周报/ 子目录，历史记录带路径。"""
    monkeypatch.setattr(settings, "APP_MODE", "local")  # 桌面态（owner 恒为 local）

    r = await client.put("/api/v1/settings/export-root",
                         json={"path": str(tmp_path / "out")})
    assert r.status_code == 200

    r = await client.post("/api/v1/history/weekly-export",
                          json={"week_offset": 0, "save_path": "auto"})
    assert r.status_code == 200, r.text
    body = r.json()
    saved = Path(body["saved_to"])
    assert saved.parent == tmp_path / "out" / "周报"
    assert saved.exists() and saved.stat().st_size > 0
    assert body["save_path"] == str(saved)

    hist = (await client.get("/api/v1/resumes/export-history")).json()["items"]
    assert any(i["template"] == "weekly" and i["save_path"] == str(saved) for i in hist)


@pytest.mark.asyncio
async def test_weekly_export_web_rejects_save_path(client):
    """web 态带 save_path → 403（不允许服务端落盘，同 export-kit 惯例）。"""
    hd = _sid()
    r = await client.post("/api/v1/history/weekly-export",
                          json={"week_offset": 0, "save_path": "auto"}, headers=hd)
    assert r.status_code == 403
