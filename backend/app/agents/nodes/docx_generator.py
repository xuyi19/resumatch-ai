import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def generate_resume_docx(data: dict) -> bytes:
    """把优化后的简历字典生成 Word 文档，返回字节流"""
    doc = Document()

    # 姓名
    name = data.get("name") or "个人简历"
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(name)
    run.bold = True
    run.font.size = Pt(22)

    # 联系方式
    if data.get("contact"):
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = c.add_run(data["contact"])
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph()

    def add_section(title_text: str, lines: list):
        if not lines:
            return
        h = doc.add_paragraph()
        run = h.add_run(title_text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x14, 0xB8, 0xA6)

        for line in lines:
            p = doc.add_paragraph(line, style="List Bullet")
            p.paragraph_format.space_after = Pt(2)

        doc.add_paragraph()

    # 个人简介
    if data.get("summary"):
        add_section("个人简介", [data["summary"]])

    add_section("教育经历", data.get("education", []))
    add_section("工作经历", data.get("experience", []))
    add_section("项目经历", data.get("projects", []))

    # 技能
    skills = data.get("skills", [])
    if skills:
        h = doc.add_paragraph()
        run = h.add_run("技能")
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x14, 0xB8, 0xA6)
        p = doc.add_paragraph(" · ".join(skills))
        p.paragraph_format.space_after = Pt(2)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()