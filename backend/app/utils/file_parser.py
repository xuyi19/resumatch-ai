from pathlib import Path

from loguru import logger


# 低于该长度视为「几乎无文本」（扫描版 PDF 的典型特征）
_LOW_TEXT_THRESHOLD = 50


def parse_pdf(path: str | Path) -> tuple[str, int, int]:
    """从 PDF 提取纯文本。

    返回 (text, image_only_pages, total_pages)：
    image_only_pages 为「几乎无文本且含图片」的页数，用于识别扫描版/图片型 PDF（B2）。
    """
    import pymupdf  # PyMuPDF

    text_parts = []
    image_only_pages = 0
    total_pages = 0
    with pymupdf.open(str(path)) as doc:
        total_pages = len(doc)
        for page in doc:
            text = page.get_text("text")
            text_parts.append(text)
            if len(text.strip()) < 10 and page.get_images():
                image_only_pages += 1
    return "\n".join(text_parts).strip(), image_only_pages, total_pages


def parse_docx(path: str | Path) -> str:
    """从 DOCX 提取纯文本（含表格）"""
    from docx import Document

    doc = Document(str(path))
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def parse_resume_file(path: str | Path) -> str:
    """根据扩展名自动选择解析器"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text, image_only_pages, total_pages = parse_pdf(path)
        # B2：整本几乎无文本且主要是图片页 → 明确提示扫描版，而非笼统「解析后为空」
        if not text or (len(text) < _LOW_TEXT_THRESHOLD
                        and total_pages > 0 and image_only_pages >= max(1, total_pages - 1)):
            raise ValueError(
                "检测到扫描版/图片型 PDF，无法提取文字。"
                "请换文字版 PDF，或先完成 OCR（如 WPS/Adobe 的识别功能）后再上传"
            )
    elif suffix in (".docx", ".doc"):
        text = parse_docx(path)
    elif suffix == ".txt":
        text = path.read_text(encoding="utf-8")
    else:
        raise ValueError(f"不支持的文件类型: {suffix}（支持 pdf/docx/txt）")

    if not text:
        raise ValueError(f"文件解析后为空: {path}")

    logger.info(f"解析文件 {path.name}，提取文本 {len(text)} 字")
    return text


def parse_bytes(data: bytes, filename: str) -> str:
    """从字节流解析（用于 FastAPI UploadFile）"""
    import tempfile

    suffix = Path(filename).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(data)
        tmp_path = tmp.name

    try:
        return parse_resume_file(tmp_path)
    finally:
        Path(tmp_path).unlink(missing_ok=True)