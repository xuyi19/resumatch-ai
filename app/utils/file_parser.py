from pathlib import Path

from loguru import logger


def parse_pdf(path: str | Path) -> str:
    """从 PDF 提取纯文本"""
    import pymupdf  # PyMuPDF

    text_parts = []
    with pymupdf.open(str(path)) as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts).strip()


def parse_docx(path: str | Path) -> str:
    """从 DOCX 提取纯文本（含表格）"""
    from docx import Document

    doc = Document(str(path))
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def parse_resume_file(path: str | Path) -> str:
    """根据扩展名自动选择解析器"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = parse_pdf(path)
    elif suffix in (".docx", ".doc"):
        text = parse_docx(path)
    elif suffix == ".txt":
        text = path.read_text(encoding="utf-8")
    else:
        raise ValueError(f"不支持的文件类型: {suffix}（支持 pdf/docx/txt）")

    if not text:
        raise ValueError(f"文件解析后为空: {path}")

    logger.info(f"解析文件 {path.name}，提取文本 {len(text)} 字")
    return text


def parse_bytes(data: bytes, filename: str) -> str:
    """从字节流解析（用于 FastAPI UploadFile）"""
    import tempfile

    suffix = Path(filename).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(data)
        tmp_path = tmp.name

    try:
        return parse_resume_file(tmp_path)
    finally:
        Path(tmp_path).unlink(missing_ok=True)