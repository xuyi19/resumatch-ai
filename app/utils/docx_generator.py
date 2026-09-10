import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEMPLATES = {
    "classic": {
        "name_color": (0x1F, 0x29, 0x37),
        "section_color": (0x0D, 0x94, 0x88),
        "text_color": (0x37, 0x41, 0x51),
        "line_color": "E5E7EB",
    },
    "modern": {
        "name_color": (0x0F, 0x17, 0x2A),
        "section_color": (0x14, 0xB8, 0xA6),
        "text_color": (0x33, 0x33, 0x33),
        "line_color": "14B8A6",
    },
    "minimal": {
        "name_color": (0x00, 0x00, 0x00),
        "section_color": (0x66, 0x66, 0x66),
        "text_color": (0x33, 0x33, 0x33),
        "line_color": "EEEEEE",
    },
}


def _add_bottom_border(paragraph, color="E5E7EB", size=6):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_bullet(doc, line, colors):
    p = doc.add_paragraph(line, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    for r in p.runs:
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(*colors["text_color"])
    return p


def _build_classic(doc, data):
    colors = TEMPLATES["classic"]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data.get("name") or "个人简历")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(*colors["name_color"])

    if data.get("contact"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(data["contact"])
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    def section(title, lines):
        if not lines:
            return
        p = doc.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(*colors["section_color"])
        _add_bottom_border(p, colors["line_color"])
        for line in lines:
            _add_bullet(doc, line, colors)
        doc.add_paragraph()

    if data.get("summary"):
        section("个人简介", [data["summary"]])
    section("教育经历", data.get("education", []))
    section("工作经历", data.get("experience", []))
    section("项目经历", data.get("projects", []))

    skills = data.get("skills", [])
    if skills:
        p = doc.add_paragraph()
        r = p.add_run("技能")
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(*colors["section_color"])
        _add_bottom_border(p, colors["line_color"])
        p2 = doc.add_paragraph(" · ".join(skills))
        for r in p2.runs:
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(*colors["text_color"])


def _build_modern(doc, data):
    colors = TEMPLATES["modern"]

    p = doc.add_paragraph()
    r = p.add_run(data.get("name") or "个人简历")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(*colors["name_color"])
    _add_bottom_border(p, colors["line_color"], size=14)

    if data.get("contact"):
        p = doc.add_paragraph()
        r = p.add_run(data["contact"])
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    def section(title, lines):
        if not lines:
            return
        p = doc.add_paragraph()
        r = p.add_run("▍ ")
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(*colors["section_color"])
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(*colors["name_color"])
        for line in lines:
            p = doc.add_paragraph(line)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(*colors["text_color"])
        doc.add_paragraph()

    if data.get("summary"):
        section("个人简介", [data["summary"]])
    section("教育经历", data.get("education", []))
    section("工作经历", data.get("experience", []))
    section("项目经历", data.get("projects", []))

    skills = data.get("skills", [])
    if skills:
        p = doc.add_paragraph()
        r = p.add_run("▍ ")
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(*colors["section_color"])
        r = p.add_run("技能")
        r.bold = True
        r.font.size = Pt(13)
        p2 = doc.add_paragraph(" / ".join(skills))
        p2.paragraph_format.left_indent = Cm(0.5)
        for r in p2.runs:
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(*colors["text_color"])


def _build_minimal(doc, data):
    colors = TEMPLATES["minimal"]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data.get("name") or "个人简历")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(*colors["name_color"])

    if data.get("contact"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(data["contact"])
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.add_paragraph()

    def section(title, lines):
        if not lines:
            return
        p = doc.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(*colors["section_color"])
        for line in lines:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(*colors["text_color"])
        doc.add_paragraph()

    if data.get("summary"):
        section("个人简介", [data["summary"]])
    section("教育经历", data.get("education", []))
    section("工作经历", data.get("experience", []))
    section("项目经历", data.get("projects", []))

    skills = data.get("skills", [])
    if skills:
        p = doc.add_paragraph()
        r = p.add_run("技能")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(*colors["section_color"])
        p2 = doc.add_paragraph(" · ".join(skills))
        for r in p2.runs:
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(*colors["text_color"])


def generate_resume_docx(data: dict, template: str = "classic") -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    builder = {
        "classic": _build_classic,
        "modern": _build_modern,
        "minimal": _build_minimal,
    }.get(template, _build_classic)

    builder(doc, data)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()